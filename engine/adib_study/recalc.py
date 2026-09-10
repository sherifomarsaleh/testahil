#!/usr/bin/env python3
"""Recalculate the DELIVERED ADIB-Egypt workbook and reconcile it against the model.

Three gates, in increasing strength:

  1. EVERY formula in the delivered file must evaluate. Anything the evaluator cannot
     parse is a FAILURE, never a skip — a formula nobody can evaluate is a number nobody
     can check.
  2. Every projected line on the Income Statement and Balance Sheet must reproduce what
     compute.py computed for that year. This is the gate that makes a formula-driven
     workbook safe: a formula that computes the right thing the wrong way, or points one
     row off, fails HERE rather than shipping a different number from the study.
  3. The three present-value lenses and the central must reproduce study_numbers.json.

The evaluator is an independent reimplementation rather than the library that wrote the
file, and it runs on the DELIVERED workbook, not on the builder.
"""
import json, os, sys
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import openpyxl
import xlcalc

# THE RECALCULATOR OPENS THE WORKBOOK THIS STUDY DELIVERS, NOT A NAME TYPED HERE.
# It named ADIB_Valuation_Model_09092026.xlsx — the 9-SEPTEMBER edition — and went on
# naming it while the builder wrote the 10-September one beside it. So the one check whose
# whole job is to prove the delivered workbook reproduces the model was reconciling
# YESTERDAY'S workbook and reporting its agreement as this edition's. It passed today while
# the delivered model published a cost of equity 52bp adrift of the study. That is [L-066]
# exactly, and it is the reason edition.py exists: the edition is written once and every
# artefact name derives from it. The same defect, in the same line, was found in SWDY's
# recalculator on 09-09-2026 — one fact, two files [R-ENF-03].
import edition as _ed
XLSX = os.path.join(HERE, _ed.MODEL_XLSX)
wb = openpyxl.load_workbook(XLSX)
D = json.load(open(os.path.join(HERE, 'study_numbers.json'), encoding='utf-8'))
P5, LN = D['projection'], D['lenses']
BK = xlcalc.Book(wb)

TOL = 0.005          # half a per cent — a formula is either right or it is not
problems = []

# ---- gate 1 -----------------------------------------------------------------------
nform, dead = 0, []
for sh, coord in BK.formula_cells():
    nform += 1
    try:
        BK.cell_value(sh, coord)
    except Exception as ex:
        dead.append('%s!%s: %s' % (sh, coord, ex))
print('gate 1 — formulas found %d, unresolvable %d' % (nform, len(dead)))
for e in dead[:20]:
    print('    ', e)
problems += ['unresolvable %s' % e for e in dead]


def val(sheet, coord):
    v = BK.cell_value(sheet, coord)
    return float(v) if v is not None else None


def near(a, b, tol=TOL):
    if a is None or b is None:
        return False
    if abs(b) < 1e-9:
        return abs(a) < 1e-6
    return abs(a / b - 1) <= tol


COLS = ['E', 'F', 'G', 'H', 'I']       # the five forecast years on Income Statement
BCOL = ['C', 'D', 'E', 'F', 'G']       # the same years on Balance Sheet


def row_of(sheet, label):
    """Find a row by its LABEL, never by a hard-coded index.

    A recalculation pinned to row numbers is a recalculation that goes quietly wrong the
    first time a row is inserted above it — and it agrees with the builder for the wrong
    reason, because both were edited by the same hand on the same day.
    """
    ws = wb[sheet]
    for r in range(1, ws.max_row + 1):
        if str(ws.cell(row=r, column=1).value or '').strip() == label:
            return r
    raise KeyError('%s!%s not found' % (sheet, label))


IS_ROWS = {'fin_income': row_of('Income Statement', 'Financing income'),
           'cost_funds': row_of('Income Statement', 'Cost of deposits'),
           'net_funds': row_of('Income Statement', 'Net income from funds'),
           'net_fees': row_of('Income Statement', 'Net fees and commissions'),
           'other_nii': row_of('Income Statement', 'Other non-interest income'),
           'admin': row_of('Income Statement', 'Administrative expenses'),
           'other_op': row_of('Income Statement', 'Other operating expenses'),
           'ecl': row_of('Income Statement', 'Expected credit losses'),
           'pbt': row_of('Income Statement', 'Profit before tax'),
           'tax': row_of('Income Statement', 'Tax'),
           'np': row_of('Income Statement', 'Net profit'),
           'np_parent': row_of('Income Statement', 'Attributable to the bank')}
BS_ROWS = {'financing': row_of('Balance Sheet', 'Financing to customers, closing'),
           'total_assets': row_of('Balance Sheet', 'Total assets, closing'),
           'avg_total_assets': row_of('Balance Sheet', 'Average total assets'),
           'equity': row_of('Balance Sheet', 'Attributable equity, closing'),
           'dividend': row_of('Balance Sheet', 'Dividend'),
           'fcfe': row_of('Balance Sheet', 'Net flow to shareholders')}
RAT_ROWS = {'nim': row_of('Income Statement', 'Net interest margin'),
            'cost_income': row_of('Income Statement', 'Cost-to-income'),
            'roa': row_of('Income Statement', 'Return on average assets'),
            'roe': row_of('Income Statement', 'Return on average equity'),
            'eps': row_of('Income Statement', 'Earnings per share'),
            'bvps': row_of('Income Statement', 'Book value per share')}

# ---- gate 2 -----------------------------------------------------------------------
n2 = 0
for j, p in enumerate(P5):
    for k, row in IS_ROWS.items():
        got = val('Income Statement', '%s%d' % (COLS[j], row))
        n2 += 1
        if not near(got, p[k]):
            problems.append('Income Statement!%s%d (%s %d): workbook %s vs model %s'
                            % (COLS[j], row, k, p['year'], got, p[k]))
    for k, row in BS_ROWS.items():
        got = val('Balance Sheet', '%s%d' % (BCOL[j], row))
        n2 += 1
        if not near(got, p[k]):
            problems.append('Balance Sheet!%s%d (%s %d): workbook %s vs model %s'
                            % (BCOL[j], row, k, p['year'], got, p[k]))
    for k, row in RAT_ROWS.items():
        got = val('Income Statement', '%s%d' % (COLS[j], row))
        n2 += 1
        if not near(got, p[k]):
            problems.append('Income Statement!%s%d (%s %d): workbook %s vs model %s'
                            % (COLS[j], row, k, p['year'], got, p[k]))
print('gate 2 — projected cells reconciled %d' % n2)

# ---- gate 3 -----------------------------------------------------------------------
FV = 'Fundamental Valuation'
HEAD = [(FV, 'B%d' % row_of(FV, 'DIVIDEND DISCOUNT — value per share  (THE CENTRAL)'),
         LN['dividend_discount'], 'dividend discount'),
        (FV, 'B%d' % row_of(FV, 'Free cash flow to equity — value per share'),
         LN['free_cash_flow_to_equity'], 'free cash flow to equity'),
        (FV, 'B%d' % row_of(FV, 'Residual income — value per share'),
         LN['residual_income'], 'residual income'),
        ('Summary', 'B%d' % row_of('Summary', 'CENTRAL  (dividend discount)'),
         D['central'], 'the central on Summary'),
        ('Summary', 'B%d' % row_of('Summary', 'Low  (free cash flow to equity)'),
         D['fair']['bear'], 'the low end of the envelope'),
        ('Summary', 'B%d' % row_of('Summary', 'High  (residual income)'),
         D['fair']['full'], 'the high end of the envelope'),
        ('Summary', 'B%d' % row_of('Summary', 'Market capitalisation, EGP million'),
         D['meta']['mktcap'], 'market capitalisation')]
for sh, coord, want, lab in HEAD:
    got = val(sh, coord)
    if not near(got, want):
        problems.append('%s!%s (%s): workbook %s vs study %s' % (sh, coord, lab, got, want))
print('gate 3 — headline reconciliations %d' % len(HEAD))

print()
if problems:
    print('RESULT: FAIL — %d problem(s)' % len(problems))
    for p_ in problems[:40]:
        print('   ', p_)
    sys.exit(1)
print('RESULT: PASS — %d formulas evaluated, %d projected cells and %d headline figures '
      'reconciled, zero mismatches' % (nform, n2, len(HEAD)))

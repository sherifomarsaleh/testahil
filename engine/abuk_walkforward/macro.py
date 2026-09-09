"""ABUK fundamental walk-forward — exogenous macro and commodity record.

Tier C throughout, and used ONLY for variables OUTSIDE the company: Egypt
inflation, the pound against the dollar, and the world urea price.  No ABUK
figure comes from here; every company number in panel.json is tier A off the
issuer's own statements.

ABUK's fiscal year runs 1 July to 30 June through FY-Jun-2025, so a calendar
series is converted as the mean of the two calendar years the fiscal year spans
and a monthly series is averaged over the twelve months it actually covers.
Both conversions are DERIVED and marked as such.
"""
import io, json, os, subprocess, datetime

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "macro.json")

WDI = ("https://api.worldbank.org/v2/country/EGY/indicator/%s"
       "?format=json&per_page=300&date=2004:2026")
CMO = ("https://thedocs.worldbank.org/en/doc/"
       "74e8be41ceb20fa0da750cda2f6b9e4e-0050012026/related/"
       "CMO-Historical-Data-Monthly.xlsx")


def _get(url, binary=False):
    p = subprocess.run(["curl", "-sS", "-L", "--max-time", "180", url],
                       capture_output=True)
    if p.returncode:
        raise SystemExit("fetch failed %s: %s" % (url, p.stderr[:200]))
    return p.stdout if binary else p.stdout.decode("utf8", "replace")


def wdi(code):
    js = json.loads(_get(WDI % code))
    return {r["date"]: r["value"] for r in js[1] if r["value"] is not None}


def cmo_urea():
    import openpyxl
    raw = _get(CMO, binary=True)
    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
    ws = wb["Monthly Prices"]
    hdr_row = col = None
    for r in range(1, 12):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str) and "urea" in v.lower():
                hdr_row, col = r, c
    if col is None:
        raise SystemExit("urea column not found in the CMO monthly sheet")
    label = ws.cell(hdr_row, col).value
    out = {}
    for r in range(hdr_row + 1, ws.max_row + 1):
        k = ws.cell(r, 1).value
        v = ws.cell(r, col).value
        if not isinstance(k, str) or not isinstance(v, (int, float)):
            continue
        k = k.strip().replace("M", "-")
        if len(k) == 7 and k[:4].isdigit():
            out[k] = float(v)
    return label, out


def fy_from_monthly(monthly, fy_end_year):
    """Mean of the twelve months July (fy-1) .. June (fy). DERIVED."""
    months = [("%d-%02d" % (fy_end_year - 1, m)) for m in range(7, 13)] + \
             [("%d-%02d" % (fy_end_year, m)) for m in range(1, 7)]
    vals = [monthly[m] for m in months if m in monthly]
    return (sum(vals) / len(vals), len(vals)) if vals else (None, 0)


def fy_from_cy(cy, fy_end_year):
    """Mean of the two calendar years the fiscal year spans. DERIVED."""
    a, b = cy.get(str(fy_end_year - 1)), cy.get(str(fy_end_year))
    if a is None or b is None:
        return None
    return (a + b) / 2.0


if __name__ == "__main__":
    cpi = wdi("FP.CPI.TOTL.ZG")
    fx = wdi("PA.NUS.FCRF")
    label, urea = cmo_urea()
    doc = {
        "_retrieved": datetime.date.today().isoformat(),
        "_tier": "C",
        "_note": "World Bank WDI and the World Bank Commodity Markets monthly "
                 "pink sheet. Exogenous country and commodity variables only. "
                 "No ABUK figure is sourced here.",
        "_fy_convention": "ABUK's fiscal year is 1 July to 30 June through "
                          "FY-Jun-2025. Calendar series -> mean of the two "
                          "calendar years spanned; monthly series -> mean of "
                          "the twelve months Jul..Jun. Both DERIVED.",
        "sources": {
            "cpi_eg_pct": {"url": WDI % "FP.CPI.TOTL.ZG", "basis": "calendar year"},
            "egp_usd": {"url": WDI % "PA.NUS.FCRF", "basis": "calendar year"},
            "urea_usd_t": {"url": CMO, "series": label, "basis": "monthly"},
        },
        "cpi_eg_pct_cy": cpi,
        "egp_usd_cy": fx,
        "urea_usd_monthly": urea,
        "fiscal_year_derived": {},
    }
    for y in range(2010, 2027):
        u, n = fy_from_monthly(urea, y)
        doc["fiscal_year_derived"]["FY%d" % y] = {
            "urea_usd_t": u, "urea_months": n,
            "cpi_eg_pct": fy_from_cy(cpi, y),
            "egp_usd": fy_from_cy(fx, y),
        }
    json.dump(doc, open(OUT, "w"), indent=1, sort_keys=True)
    for y in range(2014, 2027):
        d = doc["fiscal_year_derived"]["FY%d" % y]
        print("FY%d urea %s (%d mo)  cpi %s  fx %s" % (
            y, None if d["urea_usd_t"] is None else round(d["urea_usd_t"], 1),
            d["urea_months"],
            None if d["cpi_eg_pct"] is None else round(d["cpi_eg_pct"], 2),
            None if d["egp_usd"] is None else round(d["egp_usd"], 3)))

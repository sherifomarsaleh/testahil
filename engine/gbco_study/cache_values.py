#!/usr/bin/env python3
"""Recalculate the delivered workbook so it ships with CACHED VALUES [audit finding 19].

THE DEFECT. openpyxl writes formulas and no results. A reader opening the delivered file
in any viewer that does not itself recalculate -- a browser preview, a phone, a diff tool,
most Python readers with data_only=True -- sees every formula cell EMPTY. The workbook
exists so that a reader can check the arithmetic; a file whose every computed cell reads
blank defeats the check it was built to enable. Measured on the 07-09-2026 edition: 916
cells carried a value and all 916 were literal inputs, against 851 formulas carrying none.

WHAT THIS DOES. It hands the file to LibreOffice headless, which opens it, computes every
formula and writes the results back into the same OOXML structure. Formula text, sheet
names and sheet order are asserted unchanged either side, because a recalculation that
quietly rewrites the model is worse than one that never ran.

IT IS NOT A CHECK AND DOES NOT REPLACE ONE. recalc.py is the independent evaluator that
reconciles this workbook against the study's committed numbers; this step only makes the
file readable. Running it does not tell you the model is right.
"""
import os
import shutil
import subprocess
import sys
import tempfile

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))


def _shape(path):
    wb = load_workbook(path)
    formulas = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith('='):
                    formulas['%s!%s' % (ws.title, c.coordinate)] = c.value
    return wb.sheetnames, formulas


def cache(path):
    names_before, f_before = _shape(path)
    with tempfile.TemporaryDirectory() as tmp:
        stage = os.path.join(tmp, os.path.basename(path))
        shutil.copy(path, stage)
        out = os.path.join(tmp, 'out')
        os.makedirs(out, exist_ok=True)
        proc = subprocess.run(
            ['soffice', '--headless', '--norestore', '--convert-to', 'xlsx',
             '--outdir', out, stage],
            capture_output=True, text=True, timeout=900)
        produced = os.path.join(out, os.path.basename(path))
        if not os.path.exists(produced):
            raise SystemExit('recalculation produced no file:\n%s\n%s'
                             % (proc.stdout, proc.stderr))
        names_after, f_after = _shape(produced)
        if names_after != names_before:
            raise SystemExit('the recalculation changed the sheet list:\n  %s\n  %s'
                             % (names_before, names_after))
        if f_after != f_before:
            moved = sorted(set(f_before) ^ set(f_after))
            changed = [k for k in set(f_before) & set(f_after) if f_before[k] != f_after[k]]
            raise SystemExit('the recalculation changed the model: %d cells appeared or '
                             'vanished, %d formulas rewritten. First few: %s %s'
                             % (len(moved), len(changed), moved[:5], changed[:5]))
        wbv = load_workbook(produced, data_only=True)
        cached = sum(1 for ws in wbv.worksheets for row in ws.iter_rows()
                     for c in row if c.value is not None)
        shutil.copy(produced, path)
    return len(f_before), cached


if __name__ == '__main__':
    target = sys.argv[1] if len(sys.argv) > 1 else None
    if target is None:
        cands = sorted(f for f in os.listdir(HERE)
                       if f.startswith('GBCO_Valuation_Model_') and f.endswith('.xlsx'))
        if not cands:
            raise SystemExit('no workbook found in %s' % HERE)
        target = os.path.join(HERE, cands[-1])
    n_f, n_v = cache(target)
    print('cached %s — %d formulas preserved, %d cells now carry a value'
          % (os.path.basename(target), n_f, n_v))

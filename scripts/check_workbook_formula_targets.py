#!/usr/bin/env python3
"""No formula in a delivered workbook may point at an EMPTY cell.

WHY THIS EXISTS
    TMGH's delivered workbook discounted at zero. `DCF!B9:K9` read
    `=1/(1+'Summary'!$B$14)^n`, and `Summary!B14` is blank — the discount-rate rows
    sit at 15 to 18, and the builder had a row number typed into it that stopped being
    true when a line moved. Every discount factor evaluated to 1.0000, the sum of the
    explicit years printed 758,476 against the model's 213,065, and the enterprise value
    printed 818,577 against the study's 273,166. Three times the answer, in the file a
    reader opens, under a document that published the right number.

    NOTHING IN THE BOOK WAS LOOKING. The study's own `recalc.py` reported "85 independent
    recalculations, 0 mismatches" the whole time, because it reads the workbook with
    `data_only=False` and recomputes each relationship from the model — it never evaluates
    a single one of the workbook's own formulas. `check_workbook_values.py` RUNS that
    recalculation, and deliberately does not prescribe what it must check, for good
    reasons it states. So a green recalculation and a green gate both stood over a
    workbook that was wrong by a factor of three, and it took a narrative audit to find.

    That is the defect class this house keeps paying for: an expensive read finds a
    mechanical fault, and nothing converts the class into a cheap permanent check.

WHAT THIS GATE REQUIRES
    Every cell reference inside every formula of a delivered workbook must resolve to a
    cell that holds something. A reference into a RANGE is exempt — an empty cell inside
    a `SUM(B11:K11)` is ordinary and means zero on purpose. A single-cell reference to a
    blank is never ordinary: it is a row number that moved, a sheet that was renamed, or
    a line that was never written.

    It is a STRUCTURAL test and needs no evaluation engine, so it costs milliseconds and
    cannot be satisfied by moving a number. It does not check that a formula is RIGHT —
    only that it reaches something. TMGH's did not.

THE POPULATION IS ANCHORED BOTH WAYS  [R-ENF-04]
    Each study directory must yield a delivered workbook: `edition.py`'s MODEL_XLSX where
    the study declares one, otherwise the newest `*_Valuation_Model_*.xlsx` it holds. A
    study directory with no workbook at all is REPORTED, not skipped. A run that examined
    zero workbooks FAILS, and so does a run that parsed zero formulas across present
    workbooks — an absent answer wearing a clean one's costume.

    Only workbooks THIS HOUSE BUILDS are held. A third-party source file in a study
    directory (a downloaded country-premium sheet) is not our artefact and is not our
    standard; the `*_Valuation_Model_*` name is what separates them.

THE RATCHET  [R-ENF-02]
    `engine/build_depth_audit/workbook_formula_targets_outstanding.json`. Entries record
    the workbook and the count of dangling references at the time of recording; the list
    may only ever SHORTEN, and an entry whose count has GROWN goes red. --prune rewrites.

USAGE
    python3 scripts/check_workbook_formula_targets.py
    python3 scripts/check_workbook_formula_targets.py --prune
"""
import glob
import io
import json
import os
import re
import sys
import warnings

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

warnings.filterwarnings('ignore')

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ENGINE = os.path.join(ROOT, 'engine')
RATCHET = os.path.join(ENGINE, 'build_depth_audit',
                       'workbook_formula_targets_outstanding.json')

# A house workbook. The name is the discriminator because a study directory may also hold
# a third-party source file, which is evidence we read, not an artefact we build.
HOUSE = '*_Valuation_Model_*.xlsx'

# A range — an empty cell inside one is ordinary and deliberate.
# A BARE SHEET NAME MAY NOT START MID-EXPRESSION, AND THE HYPHEN IS WHY.
# The unquoted sheet-name class carried '-', so in "=B11-Assumptions!$B$26*E5" the
# pattern matched "B11-Assumptions" as ONE sheet name -- which resolves to no sheet, so
# the reference was dropped AND B11 was lost with it. SWDY's corporate cost load was
# reported as read by nothing while five formulas read it. The lookbehind requires a real
# boundary before an unquoted name, and a sheet whose name genuinely contains a hyphen is
# written quoted, which the first alternative still accepts.
RANGE = re.compile(r"\$?[A-Z]{1,3}\$?\d+\s*:\s*\$?[A-Z]{1,3}\$?\d+")
# A single-cell reference, with or without a sheet qualifier.
REF = re.compile(r"(?:'([^']+)'!|(?<![\w$.])([A-Za-z][\w .&]*)!)?\$?([A-Z]{1,3})\$?(\d+)\b")
# A string literal inside a formula — its text is not a reference.
STRLIT = re.compile(r'"[^"]*"')
DATE_IN_NAME = re.compile(r'_(\d{2})(\d{2})(\d{4})')


def delivered(sdir):
    """The workbook a reader opens, and why we say so."""
    ed = os.path.join(sdir, 'edition.py')
    if os.path.exists(ed):
        m = re.search(r"^MODEL_XLSX\s*=\s*(.+)$", io.open(ed, encoding='utf-8').read(), re.M)
        if m:
            ns = {}
            try:
                exec(compile(io.open(ed, encoding='utf-8').read(), ed, 'exec'), ns)
                name = ns.get('MODEL_XLSX')
            except Exception:                                          # noqa: BLE001
                name = None
            if name:
                p = os.path.join(sdir, name)
                if os.path.exists(p):
                    return p, 'edition.py declares it'
                return None, ('edition.py declares MODEL_XLSX = %r and that file is not on '
                              'disk' % name)
    cands = sorted(glob.glob(os.path.join(sdir, HOUSE)))
    where = 'newest edition on disk'
    if not cands:
        # A STUDY MAY NOT HOLD ITS OWN DELIVERED FILE. XPT's builder writes to the working
        # directory and the published workbook lives in files/ — which is where the reader
        # actually gets it, so that is where this looks next rather than reporting an
        # absence that is really a location.
        tk = os.path.basename(sdir)[:-len('_study')].lower()
        cands = sorted(p for p in glob.glob(os.path.join(ROOT, 'files', HOUSE))
                       if os.path.basename(p).lower().startswith(tk))
        where = 'published copy in files/'
    if not cands:
        return None, 'no %s in the directory or in files/' % HOUSE

    def key(p):
        m = DATE_IN_NAME.search(os.path.basename(p))
        return (m.group(3), m.group(2), m.group(1)) if m else ('0000', '00', '00')

    return max(cands, key=key), where


def dangling(path):
    """Every single-cell reference in the workbook that lands on a blank cell."""
    wb = load_workbook(path, data_only=False)
    names = set(wb.sheetnames)
    hits, parsed = [], 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if not (isinstance(v, str) and v.startswith('=')):
                    continue
                parsed += 1
                body = RANGE.sub(' ', STRLIT.sub(' ', v))
                for m in REF.finditer(body):
                    sheet = m.group(1) or m.group(2)
                    if sheet is not None and sheet not in names:
                        continue          # a function name, or a book we cannot resolve
                    tgt = wb[sheet] if sheet else ws
                    try:
                        tv = tgt.cell(row=int(m.group(4)),
                                      column=column_index_from_string(m.group(3))).value
                    except Exception:                                   # noqa: BLE001
                        continue
                    if tv is None or (isinstance(tv, str) and not tv.strip()):
                        hits.append('%s!%s  %s  -> empty %s!%s%s'
                                    % (ws.title, c.coordinate, v[:56],
                                       sheet or ws.title, m.group(3), m.group(4)))
    return hits, parsed


def load_ratchet():
    if not os.path.exists(RATCHET):
        return {}
    return json.load(open(RATCHET)).get('dangling', {})


def main():
    prune = '--prune' in sys.argv
    ratchet = load_ratchet()
    sdirs = sorted(glob.glob(os.path.join(ENGINE, '*_study')))
    if not sdirs:
        print('FAIL — no study directories resolved; the population is empty [R-ENF-04]')
        return 1

    fails, examined, formulas, keep = [], 0, 0, {}
    for sdir in sdirs:
        tk = os.path.basename(sdir)[:-len('_study')]
        path, why = delivered(sdir)
        if path is None:
            fails.append('%-14s NO DELIVERED WORKBOOK — %s' % (tk, why))
            continue
        try:
            hits, parsed = dangling(path)
        except Exception as e:                                          # noqa: BLE001
            fails.append('%-14s UNREADABLE — %s (%s)' % (tk, e, os.path.basename(path)))
            continue
        examined += 1
        formulas += parsed
        rec = ratchet.get(tk)
        if not hits:
            if rec is not None and not prune:
                fails.append('%-14s is CLEAN and still on the ratchet — remove it '
                             '(--prune)' % tk)
            continue
        if rec is None:
            fails.append('%-14s %d dangling reference(s) in %s\n%s'
                         % (tk, len(hits), os.path.basename(path),
                            '\n'.join('                 ' + h for h in hits[:8])))
        elif len(hits) > rec.get('count', 0):
            fails.append('%-14s WORSENED — %d dangling references against %d recorded'
                         % (tk, len(hits), rec.get('count', 0)))
        else:
            keep[tk] = {'workbook': os.path.basename(path), 'count': len(hits),
                        'note': rec.get('note', '')}

    if examined == 0:
        print('FAIL — zero workbooks examined across %d study directories [R-ENF-04]'
              % len(sdirs))
        return 1
    if formulas == 0:
        print('FAIL — %d workbooks opened and not one formula parsed; an empty result is '
              'not a clean result [R-ENF-04]' % examined)
        return 1

    if prune:
        os.makedirs(os.path.dirname(RATCHET), exist_ok=True)
        json.dump({'dangling': keep}, open(RATCHET, 'w'), indent=1, sort_keys=True)
        print('ratchet rewritten: %d entr%s' % (len(keep), 'y' if len(keep) == 1 else 'ies'))

    print('%d delivered workbooks examined, %d formula cells parsed, %d excused by the '
          'ratchet' % (examined, formulas, len(keep)))
    if fails:
        print('\nFAIL')
        for f in fails:
            print('   ' + f)
        return 1
    print('OK — every formula reference in every delivered workbook lands on a cell that '
          'holds something.')
    return 0


if __name__ == '__main__':
    sys.exit(main())

#!/usr/bin/env python3
"""Render a study's MARKET_DISSENT markdown to the delivered .docx, then to PDF.

WHY THIS EXISTS. [R-GAP-02] makes a study past the publication limit write its case
before it may publish, and the case is the READER'S, not the file system's. SWDY
published on 14-09-2026 with MARKET_DISSENT_14-09-2026.md filed, gating the publish
exactly as the rule intends — and with nothing on the ticker page pointing at it. The
document existed, was public in the repository, and was unreachable from the one page
anybody opens.

TWO DEFECTS IN THE FIRST VERSION OF THIS BUILDER, both found by reading the PDF it
produced rather than by any check here:

  TEXT BROKE MID-SENTENCE. Markdown is hard-wrapped at the column, and this emitted
  one PARAGRAPH PER LINE, so every wrapped sentence became three ragged paragraphs.
  A block now accumulates until something ends it — a blank line, a heading, a list
  marker, a table row, a rule — and is written as one paragraph.

  NO HOUSE STYLE AT ALL. python-docx defaults are not a brand. The palette is read
  from the site's own CSS custom properties and the delivered study's own styles:
  ink #0E2726, gold #C0A45F, muted #5B7270, rules #D9E4E2, headings in the deep
  green 1C3A36 the study itself uses, all set in Calibri to match.

The five headings [R-GAP-02] requires, and the DISSENT_AT_GAP line the gate parses,
are carried through exactly as filed. This changes presentation and never content.

    python3 scripts/build_dissent_docx.py SWDY
"""
from __future__ import annotations

import glob
import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "engine"))
import make_pdf  # noqa: E402  — the repo's own renderer, never a second copy

# THE PALETTE IS NOT INVENTED HERE. --ink/--gold/--muted/--line are the site's own
# CSS custom properties in assets/, and 1C3A36 is the heading colour already in the
# delivered study's styles.xml. A document that carries a different green from the
# page it is downloaded from reads as somebody else's document.
INK = RGBColor(0x0E, 0x27, 0x26)
HEAD = RGBColor(0x1C, 0x3A, 0x36)
GOLD = RGBColor(0xC0, 0xA4, 0x5F)
MUTED = RGBColor(0x5B, 0x72, 0x70)
LINE = "D9E4E2"
BAND = "F6F8F7"


def _shade(cell, hexfill: str) -> None:
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hexfill)
    cell._tc.get_or_add_tcPr().append(el)


def _repeat_header(row) -> None:
    """Mark a table row as a header that REPEATS on every page it continues onto.

    Without this a table that crosses a page break leaves its first visible row
    orphaned — page 2 of the first build opened with a bare "(adopted 10.34%)" and
    no column above it to say what the number was. A reader cannot read a column
    whose heading is on the previous page.
    """
    pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    pr.append(el)


def _no_split(row) -> None:
    """Keep a row's cells on one page rather than tearing a cell in half."""
    pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "true")
    pr.append(el)


def _borders(table) -> None:
    tbl = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:color"), LINE)
        borders.append(e)
    tbl.append(borders)


def _runs(par, text: str, colour=INK, size=10.5) -> None:
    """Write text into a paragraph, honouring **bold**, *italic* and `code`."""
    for chunk in re.split(r"(\*\*[^*]+\*\*|`[^`]+`|(?<!\*)\*[^*]+\*(?!\*))", text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**"):
            r = par.add_run(chunk[2:-2])
            r.bold = True
        elif chunk.startswith("`") and chunk.endswith("`"):
            r = par.add_run(chunk[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(size - 1)
        elif chunk.startswith("*") and chunk.endswith("*"):
            r = par.add_run(chunk[1:-1])
            r.italic = True
        else:
            r = par.add_run(chunk)
        r.font.color.rgb = colour
        r.font.size = Pt(size)


def build(ticker: str) -> str:
    tk = ticker.upper()
    sdir = os.path.join(ROOT, "engine", f"{tk.lower()}_study")
    hits = sorted(glob.glob(os.path.join(sdir, "MARKET_DISSENT_*.md")))
    if not hits:
        sys.exit(f"FAIL: {tk} has no MARKET_DISSENT_*.md in {sdir}")
    src = hits[-1]
    stamp = re.search(r"MARKET_DISSENT_(\d{2}-\d{2}-\d{4})\.md$", os.path.basename(src))
    if not stamp:
        sys.exit(f"FAIL: {os.path.basename(src)} carries no DD-MM-YYYY stamp")

    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(0.9)
        s.top_margin = s.bottom_margin = Inches(0.8)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12

    buf: list[str] = []
    rows: list[list[str]] = []

    def flush_text() -> None:
        if not buf:
            return
        text = " ".join(x.strip() for x in buf).strip()
        buf.clear()
        if not text:
            return
        # The gate parses this line; it is set apart so a reader sees it too.
        if re.match(r"^DISSENT[ _]AT[ _]GAP", text, re.I):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            _runs(p, text, colour=GOLD, size=12)
            for r in p.runs:
                r.bold = True
            return
        if re.match(r"^[-*] ", text):
            _runs(doc.add_paragraph(style="List Bullet"), text[2:].strip())
        elif re.match(r"^\d+\. ", text):
            _runs(doc.add_paragraph(style="List Number"), re.sub(r"^\d+\.\s*", "", text))
        else:
            _runs(doc.add_paragraph(), text)

    def flush_table() -> None:
        if not rows:
            return
        body = [r for r in rows if not re.match(r"^[\s|:-]+$", "|".join(r))]
        rows.clear()
        if not body:
            return
        width = max(len(r) for r in body)
        t = doc.add_table(rows=len(body), cols=width)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        _borders(t)
        for i, row in enumerate(body):
            _no_split(t.rows[i])
            if i == 0:
                _repeat_header(t.rows[i])
            for j in range(width):
                cell = t.rows[i].cells[j]
                cell.paragraphs[0].paragraph_format.space_after = Pt(3)
                _runs(cell.paragraphs[0],
                      row[j].strip() if j < len(row) else "",
                      colour=HEAD if i == 0 else INK, size=9.5)
                if i == 0:
                    _shade(cell, BAND)
                    for r in cell.paragraphs[0].runs:
                        r.bold = True
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def heading(text: str, level: int) -> None:
        flush_text()
        flush_table()
        h = doc.add_heading("", level=level)
        h.paragraph_format.space_before = Pt(14 if level <= 2 else 10)
        h.paragraph_format.space_after = Pt(5)
        _runs(h, text, colour=HEAD, size={1: 19, 2: 13.5, 3: 11.5}.get(level, 11))
        for r in h.runs:
            r.bold = True
        if level == 1:
            h.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for raw in open(src, encoding="utf-8").read().splitlines():
        line = raw.rstrip()
        stripped = line.strip()

        if stripped.startswith("|"):
            flush_text()
            rows.append(stripped.strip("|").split("|"))
            continue
        flush_table()

        if not stripped:
            flush_text()
            continue
        if stripped == "---":
            flush_text()
            continue
        if stripped.startswith("### "):
            heading(stripped[4:].strip(), 3)
            continue
        if stripped.startswith("## "):
            heading(stripped[3:].strip(), 2)
            continue
        if stripped.startswith("# "):
            heading(stripped[2:].strip(), 1)
            continue
        # A LIST MARKER ENDS THE PREVIOUS BLOCK but its own continuation lines do not:
        # the dissent wraps its bullets across three lines and they are one bullet.
        if re.match(r"^([-*] |\d+\. )", stripped):
            flush_text()
        buf.append(stripped)

    flush_text()
    flush_table()

    out = os.path.join(ROOT, "files", f"{tk}_Market_Dissent_{stamp.group(1)}.docx")
    doc.save(out)
    pdf = make_pdf.convert(out, os.path.join(ROOT, "files"))
    print(f"  built {os.path.relpath(out, ROOT)}")
    print(f"  built {os.path.relpath(pdf, ROOT)}")
    return pdf


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    build(sys.argv[1])

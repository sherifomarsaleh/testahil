#!/usr/bin/env python3
"""A check nobody has seen fail is not evidence.

Every condition is injected into a SANDBOX workbook and the gate must go red on it, or
stay green where the cell really is read or really is declared a reference.

THE CASE THAT MATTERS IS 4 — the declared display row. A sheet may legitimately print a
figure for context: a prior edition's rate, a basis the model does not adopt. That
exemption is the one route back to the blindness this gate closes, so it is bounded here:
the label must SAY SO ON THE PAGE, and a dead cell under an ordinary label is red however
innocent it looks.
"""
import json
import os
import shutil
import subprocess
import sys
import tempfile

from openpyxl import Workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
GATE = os.path.join('scripts', 'check_workbook_live_inputs.py')
RAT = os.path.join('engine', 'build_depth_audit', 'live_inputs_outstanding.json')


def book(path, *, dead=False, label='Cost of capital', in_range=False, other_sheet=False):
    wb = Workbook()
    a = wb.active
    a.title = 'Assumptions'
    a['A1'] = 'Assumptions — every blue cell is an input'
    a['A2'] = 'Cash conversion'
    a['B2'] = 0.087
    a['A3'] = label
    a['B3'] = 0.25
    a['A4'] = 'Shares in issue (mn)'
    a['B4'] = 1000.0
    d = wb.create_sheet('DCF')
    d['A1'] = 'Free cash flow'
    d['B1'] = 5000.0
    d['A2'] = 'Operating cash'
    d['B2'] = '=B1*Assumptions!B2'
    d['A3'] = 'Value'
    # B3 (the rate) is read only when `dead` is False
    d['B3'] = ('=B2/(1+Assumptions!B3)' if not dead else '=B2/(1+0.25)')
    d['A4'] = 'Per share'
    d['B4'] = '=B3/Assumptions!B4'
    if in_range:
        # read through a RANGE rather than a single reference
        d['B3'] = '=B2/(1+SUM(Assumptions!B3:B3))'
    if other_sheet:
        e = wb.create_sheet('Bridge')
        e['A1'] = 'Rate'
        e['B1'] = '=Assumptions!B3'
    wb.save(path)


def sandbox(**kw):
    t = tempfile.mkdtemp(prefix='live_nc_')
    os.makedirs(os.path.join(t, 'scripts'))
    os.makedirs(os.path.join(t, 'engine', 'build_depth_audit'))
    sd = os.path.join(t, 'engine', 'alpha_study')
    os.makedirs(sd)
    shutil.copy(os.path.join(ROOT, GATE), os.path.join(t, GATE))
    json.dump({'dead': {}}, open(os.path.join(t, RAT), 'w'))
    book(os.path.join(sd, 'ALPHA_Valuation_Model_01012026.xlsx'), **kw)
    return t, sd


def run(t):
    r = subprocess.run([sys.executable, GATE], cwd=t, stdout=subprocess.PIPE,
                       stderr=subprocess.STDOUT, text=True, timeout=300)
    return r.returncode, r.stdout


CASES = []


def case(name, expect_red, build):
    CASES.append((name, expect_red, build))


case('1. every input read by a formula', False, lambda: sandbox())
case('2. an input no formula reads — the rate hardcoded into the DCF instead', True,
     lambda: sandbox(dead=True))
case('3. read only through a RANGE, not a single reference', False,
     lambda: sandbox(in_range=True))
case('4. dead, but the label declares it a reference', False,
     lambda: sandbox(dead=True, label='Cost of capital — alternative basis, not used'))
case('5. dead under an ordinary label that merely sounds harmless', True,
     lambda: sandbox(dead=True, label='Cost of capital, published'))
case('6. read from a DIFFERENT sheet', False,
     lambda: sandbox(dead=True, other_sheet=True))


def no_workbook():
    t, sd = sandbox()
    os.remove(os.path.join(sd, 'ALPHA_Valuation_Model_01012026.xlsx'))
    return t, sd


def no_studies():
    t, sd = sandbox()
    shutil.rmtree(sd)
    return t, sd


def ratcheted():
    t, sd = sandbox(dead=True)
    json.dump({'dead': {'ALPHA': ['Assumptions!B3']}}, open(os.path.join(t, RAT), 'w'))
    return t, sd


def ratchet_wrong_cell():
    t, sd = sandbox(dead=True)
    json.dump({'dead': {'ALPHA': ['Assumptions!B99']}}, open(os.path.join(t, RAT), 'w'))
    return t, sd


case('7. a study with no delivered workbook is UNKNOWN, not clean', True, no_workbook)
case('8. no study directories at all [R-ENF-04]', True, no_studies)
case('9. the dead cell is on the ratchet', False, ratcheted)
case('10. the ratchet names a DIFFERENT cell', True, ratchet_wrong_cell)


def main():
    bad = []
    for name, expect_red, build in CASES:
        t, _sd = build()
        try:
            rc, out = run(t)
            red = rc != 0
            ok = red == expect_red
            print('%-58s %-5s %s' % (name, 'RED' if red else 'green',
                                     'ok' if ok else '*** NOT AS REQUIRED'))
            if not ok:
                bad.append((name, out))
        finally:
            shutil.rmtree(t, ignore_errors=True)
    if bad:
        print('\n%d condition(s) did not behave as required:' % len(bad))
        for n, o in bad:
            print('\n--- %s\n%s' % (n, o[:1200]))
        return 1
    print('\nall %d conditions behave as required' % len(CASES))
    return 0


if __name__ == '__main__':
    sys.exit(main())
